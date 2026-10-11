"""Étape 5a — Générer le cours PowerPoint (charte : voir theme.py).

Chaque slide de contenu : titre-message, contenu (texte, tableau, graphique natif ou schéma),
point de vigilance éventuel, encadré « À retenir », sources détaillées en pied de page.
Le rendu note les textes qui débordent malgré la réduction de police (contrôle qualité).
"""
import math
import re
from datetime import date

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import normalize
from . import theme as theme_mod

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre",
        "octobre", "novembre", "décembre"]

GLOSS_PER_PAGE = 12         # termes de glossaire par slide
SOURCES_PER_PAGE = 10       # sources par slide

W, H = 13.333, 7.5
ML = 0.6                     # marge gauche/droite
CW = W - 2 * ML              # largeur utile
BODY_Y = 1.45                # haut de la zone de contenu
RETENIR_Y = 6.0              # encadré « À retenir »
FOOT_Y = 6.95                # filet du pied de page


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace("%", "").replace(",", ".")
    s = s.replace("−", "-")
    return float(s)


class Deck:
    def __init__(self, course: dict):
        self.c = course
        self.t = theme_mod.load()
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 1                     # numéro de la slide courante
        self.cur = None                # (section, slide) en cours, pour le rapport qualité
        self.issues: dict = {}         # (i, j) -> [problèmes]
        self.page_map: dict = {}       # numéro de slide -> (i, j)
        self.refs = {b["id"]: b.get("court") or b.get("ref", "")[:40] for b in course.get("bibliographie", [])}
        self.palette = [self.t["red"], self.t["ink"], self.t["grey"], self.t["red_dark"], "E8918D", "BDBDBD"]

    def flag(self, msg):
        if self.cur is not None:
            self.issues.setdefault(self.cur, []).append(msg)

    # ------------------------------------------------------------------ primitives
    def rgb(self, key):
        return RGBColor.from_string(self.t.get(key, key))

    def rect(self, s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = self.rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        return sh

    @staticmethod
    def need(texts, w, size, spacing=1.22):
        cpl = max(6, (w - 0.15) * 72 / (size * 0.5))
        lines = sum(max(1, math.ceil(len(t) / cpl)) for t in texts)
        return lines * size * spacing / 72 + len(texts) * size * 0.4 / 72 + 0.1

    def text(self, s, x, y, w, h, paras, size=14, color="ink", bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=None, bullets=False, space_after=6, auto=True, mini=10):
        """`paras` : str, ou liste d'éléments str / list[(texte, {bold, color, size, font})]."""
        if isinstance(paras, str):
            paras = [paras]
        paras = [p for p in paras if p not in (None, "")]
        if auto and paras:
            flat = ["".join(str(r[0]) for r in p) if isinstance(p, list) else str(p) for p in paras]
            while size > mini and self.need(flat, w, size) > h:
                size -= 1
            if self.need(flat, w, size) > h * 1.08:
                self.flag(f"texte trop long pour sa zone : « {flat[0][:60]}… »")
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        for i, p in enumerate(paras or [""]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(space_after)
            runs = p if isinstance(p, list) else [(p, {})]
            if bullets:
                runs = [("•  ", {"color": "red", "bold": True})] + runs
            for txt, opt in runs:
                r = para.add_run()
                r.text = str(txt)
                f = r.font
                f.name = opt.get("font", font or self.t["font"])
                f.size = Pt(opt.get("size", size))
                f.bold = opt.get("bold", bold)
                f.color.rgb = self.rgb(opt.get("color", color))
        return tb

    def circle(self, s, x, y, d, label, fill="red", color="white", size=14):
        c = self.rect(s, x, y, d, d, fill, MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(label)
        r.font.size, r.font.bold, r.font.name = Pt(size), True, self.t["font"]
        r.font.color.rgb = self.rgb(color)
        return c

    def logo(self, s, x, y, h, white=False):
        path = self.t["logo_white"] if white else self.t["logo"]
        if path:
            s.shapes.add_picture(path, Inches(x), Inches(y), height=Inches(h))

    def arrow(self, s, x1, y1, x2, y2, color="grey", width=1.75):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = self.rgb(color)
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        return c

    def source_label(self, sources):
        out = []
        for src in sources or []:
            sid = str(src).split()[0].strip("[](),;")
            court = self.refs.get(sid)
            out.append(f"{src} – {court}" if court else str(src))
        return " ; ".join(out)

    # ------------------------------------------------------------------ cadres
    def content_slide(self, sl, title=None):
        self.n += 1
        s = self.prs.slides.add_slide(self.blank)
        self.rect(s, ML, 0.38, 0.55, 0.07, "red")
        self.text(s, ML, 0.5, CW - (1.6 if self.t["logo"] else 0), 0.85, title or sl.get("titre", ""), 24, "ink",
                  True, anchor=MSO_ANCHOR.MIDDLE, mini=15)
        self.logo(s, W - ML - 1.3, 0.45, 0.45)
        if sl.get("a_retenir"):
            self.rect(s, ML, RETENIR_Y, CW, 0.75, "soft")
            self.rect(s, ML, RETENIR_Y, 0.08, 0.75, "red")
            self.text(s, ML + 0.25, RETENIR_Y, CW - 0.4, 0.75,
                      [[("À retenir   ", {"bold": True, "color": "red"}), (sl["a_retenir"], {})]],
                      14, anchor=MSO_ANCHOR.MIDDLE, mini=10, space_after=0)
        if sl.get("vigilance"):
            y = self.body_bottom(sl) + 0.12
            self.rect(s, ML, y, CW, 0.58, "white", line="red_dark")
            self.text(s, ML + 0.2, y, CW - 0.4, 0.58,
                      [[("⚠ Point de vigilance   ", {"bold": True, "color": "red_dark"}), (sl["vigilance"], {})]],
                      12, anchor=MSO_ANCHOR.MIDDLE, mini=9, space_after=0)
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ML), Inches(FOOT_Y), Inches(W - ML), Inches(FOOT_Y))
        line.line.color.rgb = self.rgb("line")
        foot = ("Sources : " + self.source_label(sl.get("sources"))) if sl.get("sources") else self.c["titre"]
        self.text(s, ML, FOOT_Y + 0.04, CW - 0.8, 0.45, foot, 9, "grey", mini=7, space_after=0)
        self.text(s, W - ML - 0.7, FOOT_Y + 0.04, 0.7, 0.3, str(self.n), 9, "grey", align=PP_ALIGN.RIGHT,
                  auto=False, space_after=0)
        return s

    def body_bottom(self, sl):
        bottom = RETENIR_Y - 0.15 if sl.get("a_retenir") else FOOT_Y - 0.15
        return bottom - (0.7 if sl.get("vigilance") else 0)

    # ------------------------------------------------------------------ slides fixes
    def cover(self):
        s = self.prs.slides.add_slide(self.blank)
        self.rect(s, 0, 0, 0.35, H, "red")
        self.rect(s, 8.9, 0, W - 8.9, H, "panel")
        self.logo(s, 9.6, 0.7, 0.8)
        self.text(s, 1.1, 0.9, 6.0, 0.4, "COURS DE FORMATION", 13, "red", True, auto=False)
        self.text(s, 1.1, 1.5, 7.4, 2.8, self.c["titre"], 36, "ink", True, anchor=MSO_ANCHOR.BOTTOM, mini=22)
        self.rect(s, 1.1, 4.45, 1.2, 0.07, "red")
        self.text(s, 1.1, 4.7, 7.4, 1.5, self.c.get("sous_titre", ""), 16, "grey", mini=11)
        n = sum(1 for b in self.c.get("bibliographie", []) if b.get("nature", "Mémoire") == "Mémoire")
        byline = [self.t["auteur"], f"{MOIS[date.today().month - 1]} {date.today().year}",
                  f"cas concrets tirés de {n} mémoires d'actuariat" if n else ""]
        self.text(s, 1.1, 6.5, 7.6, 0.4, "  ·  ".join(b for b in byline if b), 11, "grey", auto=False)
        self.text(s, 9.6, 1.9, 3.2, 0.4, "SOMMAIRE", 12, "red", True, auto=False)
        items = [[(f"{i}  ", {"bold": True, "color": "red"}), (sec["titre"], {})]
                 for i, sec in enumerate(self.c["sections"], 1)]
        items += [[("•  ", {"bold": True, "color": "red"}), ("Glossaire et sources", {})]]
        self.text(s, 9.6, 2.4, 3.3, 4.4, items, 14, "ink", space_after=10, mini=9)

    def objectives(self):
        sections = [sec["titre"] for sec in self.c["sections"]]
        s = self.content_slide({"a_retenir": "Parcours : " + " → ".join(f"{i}. {t}" for i, t in enumerate(sections, 1))},
                               "Ce que vous saurez faire à la fin de ce cours")
        objs = self.c.get("objectifs", [])[:5]
        top, bottom = BODY_Y, RETENIR_Y - 0.15
        rowh = (bottom - top) / max(len(objs), 1)
        for i, o in enumerate(objs):
            y = top + i * rowh
            self.circle(s, ML + 0.1, y + (rowh - 0.5) / 2, 0.5, i + 1)
            self.text(s, ML + 0.85, y, CW - 1, rowh,
                      [[(o.get("verbe", "") + " ", {"bold": True, "color": "red"}), (o.get("texte", ""), {})]],
                      17, anchor=MSO_ANCHOR.MIDDLE, mini=11)

    def divider(self, i, sec):
        self.n += 1
        s = self.prs.slides.add_slide(self.blank)
        self.rect(s, 0, 0, W, H, "dark")
        self.rect(s, 0, H - 0.12, W, 0.12, "red")
        self.circle(s, 1.1, 2.9, 1.0, i, fill="red", color="white", size=28)
        self.text(s, 2.5, 2.4, 9.8, 1.9, sec["titre"], 34, "white", True, anchor=MSO_ANCHOR.MIDDLE, mini=20)
        self.text(s, 2.5, 4.35, 9.8, 1.2, sec.get("accroche", ""), 17, "line", mini=12)
        self.logo(s, W - ML - 1.4, H - 0.85, 0.55, white=True)

    def closing(self):
        pts = self.c.get("points_cles", [])
        s = self.content_slide({}, "Les points clés du cours")
        colw = 7.4
        bottom = FOOT_Y - 0.2
        rowh = (bottom - BODY_Y) / max(len(pts), 1)
        for i, p in enumerate(pts):
            y = BODY_Y + i * rowh
            self.circle(s, ML, y + (rowh - 0.42) / 2, 0.42, i + 1, size=12)
            self.text(s, ML + 0.6, y, colw - 0.6, rowh, p, 15, anchor=MSO_ANCHOR.MIDDLE, mini=10)
        x = ML + colw + 0.4
        self.rect(s, x, BODY_Y, W - ML - x, bottom - BODY_Y, "panel")
        self.rect(s, x, BODY_Y, W - ML - x, 0.07, "red")
        self.text(s, x + 0.25, BODY_Y + 0.25, W - ML - x - 0.5, 0.4, "Pour aller plus loin", 15, "red", True, auto=False)
        self.text(s, x + 0.25, BODY_Y + 0.8, W - ML - x - 0.5, bottom - BODY_Y - 1.0,
                  self.c.get("pour_aller_plus_loin", []), 13, bullets=True, mini=9)

    def paged_table(self, title, header, rows, widths, per_page, size=11):
        pages = max(1, math.ceil(len(rows) / per_page))
        for k in range(pages):
            chunk = rows[k * per_page:(k + 1) * per_page]
            t = title + (f" ({k + 1}/{pages})" if pages > 1 else "")
            s = self.content_slide({}, t)
            row_h = min(0.5, (FOOT_Y - 0.2 - BODY_Y) / (len(chunk) + 1))
            tbl = s.shapes.add_table(len(chunk) + 1, len(header), Inches(ML), Inches(BODY_Y), Inches(CW),
                                     Inches(row_h * (len(chunk) + 1))).table
            for j, w in enumerate(widths):
                tbl.columns[j].width = Inches(w)
            self._fill_table(tbl, header, chunk, size)

    def glossary(self):
        g = self.c.get("glossaire", [])
        if g:
            self.paged_table("Glossaire", ["Terme", "Définition", "Source"],
                             [[x.get("terme", ""), x.get("definition", ""), x.get("source", "")] for x in g],
                             [2.6, 7.6, CW - 10.2], GLOSS_PER_PAGE, 10)

    def bibliography(self):
        bib = self.c.get("bibliographie", [])
        self.paged_table("Toutes les sources de ce cours", ["Réf.", "Nature", "Référence", "Lien"],
                         [[b["id"], b.get("nature", ""), b["ref"], b["url"]] for b in bib],
                         [0.7, 2.2, 5.4, CW - 8.3], SOURCES_PER_PAGE, 10)

    # ------------------------------------------------------------------ blocs réutilisables
    def _fill_table(self, tbl, header, rows, size):
        for j, h in enumerate(header):
            self._cell(tbl.cell(0, j), h, size, "white", True, "red")
        for i, row in enumerate(rows, 1):
            for j in range(len(header)):
                val = row[j] if j < len(row) else ""
                self._cell(tbl.cell(i, j), val, size, "ink", j == 0, "panel" if i % 2 == 0 else "white",
                           mono=bool(re.fullmatch(r"[A-Z0-9_\.]{4,}", str(val))))

    def _cell(self, cell, txt, size, color, bold, fill, mono=False):
        cell.fill.solid()
        cell.fill.fore_color.rgb = self.rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        tf = cell.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = ""
        r = tf.paragraphs[0].add_run()
        r.text = str(txt)
        r.font.size, r.font.bold = Pt(size), bold
        r.font.name = self.t["mono"] if mono else self.t["font"]
        r.font.color.rgb = self.rgb(color)

    def card(self, s, x, y, w, h, title, body, number=None, body_size=16):
        self.rect(s, x, y, w, h, "panel")
        self.rect(s, x, y, w, 0.07, "red")
        tx = x + 0.25
        if number is not None:
            self.circle(s, x + 0.22, y + 0.28, 0.45, number, size=13)
            tx = x + 0.85
        self.text(s, tx, y + 0.22, x + w - tx - 0.2, 0.6, title, 17, "red", True, anchor=MSO_ANCHOR.MIDDLE, mini=11)
        if isinstance(body, list):
            self.text(s, x + 0.25, y + 0.95, w - 0.45, h - 1.05, body, body_size, bullets=True, space_after=9, mini=10)
        elif body:
            self.text(s, x + 0.25, y + 0.95, w - 0.45, h - 1.05, body, body_size, mini=10)

    def side_note(self, s, x, y, w, h, title, text):
        self.rect(s, x, y, w, h, "panel")
        self.rect(s, x, y, 0.07, h, "red")
        self.text(s, x + 0.25, y + 0.2, w - 0.45, 0.45, title, 14, "red", True, auto=False)
        self.text(s, x + 0.25, y + 0.7, w - 0.45, h - 0.85, text, 14, mini=10)

    # ------------------------------------------------------------------ types de slides
    def s_colonnes(self, s, sl):
        cols = sl.get("colonnes", [])[:2]
        enc = sl.get("encadre")
        bottom = self.body_bottom(sl)
        h = (bottom - BODY_Y) - (1.3 if enc else 0)
        gap = 0.35
        w = (CW - gap) / 2
        for i, col in enumerate(cols):
            self.card(s, ML + i * (w + gap), BODY_Y, w, h, col.get("titre", ""), col.get("points", []))
        if enc:
            y = BODY_Y + h + 0.2
            self.rect(s, ML, y, CW, bottom - y, "white", line="line")
            self.rect(s, ML, y, 0.08, bottom - y, "red")
            self.text(s, ML + 0.25, y + 0.05, CW - 0.4, bottom - y - 0.1,
                      [[(enc.get("titre", "") + " — ", {"bold": True, "color": "red"}), (enc.get("texte", ""), {})]],
                      14, anchor=MSO_ANCHOR.MIDDLE, mini=10)

    def s_cartes(self, s, sl):
        cards = sl.get("cartes", [])[:4]
        n = max(len(cards), 1)
        gap = 0.3
        w = (CW - gap * (n - 1)) / n
        h = self.body_bottom(sl) - BODY_Y
        for i, c in enumerate(cards):
            self.card(s, ML + i * (w + gap), BODY_Y, w, h, c.get("titre", ""), c.get("texte", ""), number=i + 1,
                      body_size=17)

    def s_etapes(self, s, sl):
        steps = sl.get("etapes", [])[:5]
        n = max(len(steps), 1)
        extra = sl.get("texte")
        bottom = self.body_bottom(sl)
        gap = 0.12
        w = (CW - gap * (n - 1)) / n
        head_h = 0.75
        box_h = (bottom - BODY_Y - head_h - 0.1) - (0.95 if extra else 0)
        for i, st in enumerate(steps):
            x = ML + i * (w + gap)
            shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            self.rect(s, x, BODY_Y, w + (0.18 if i < n - 1 else 0), head_h, "red" if i % 2 == 0 else "red_dark", shape)
            self.text(s, x + (0.15 if i == 0 else 0.35), BODY_Y, w - 0.45, head_h,
                      [[(f"{i + 1}  ", {"bold": True}), (st.get("titre", ""), {"bold": True})]],
                      14, "white", anchor=MSO_ANCHOR.MIDDLE, mini=9, space_after=0)
            self.rect(s, x, BODY_Y + head_h + 0.1, w, box_h, "panel")
            self.text(s, x + 0.12, BODY_Y + head_h + 0.2, w - 0.24, box_h - 0.2, st.get("texte", ""), 16, mini=10)
        if extra:
            self.text(s, ML, bottom - 0.85, CW, 0.85, extra, 14, mini=10)

    def s_formule(self, s, sl):
        bottom = self.body_bottom(sl)
        y = BODY_Y
        if sl.get("intro"):
            self.text(s, ML, y, CW, 0.6, sl["intro"], 15, mini=11)
            y += 0.7
        forms = sl.get("formules", [])[:3]
        n = max(len(forms), 1)
        h = (bottom - y - 0.15 * (n - 1)) / n
        for f in forms:
            self.rect(s, ML, y, CW, h, "panel")
            self.rect(s, ML, y, 0.08, h, "red")
            self.text(s, ML + 0.3, y + 0.08, CW - 0.6, 0.35, f.get("nom", ""), 13, "red", True, auto=False, space_after=0)
            self.text(s, ML + 0.3, y + 0.42, CW - 0.6, h * 0.45, f.get("expression", ""), 18, "ink",
                      font=self.t["mono"], anchor=MSO_ANCHOR.MIDDLE, mini=10, space_after=0)
            self.text(s, ML + 0.3, y + 0.42 + h * 0.45, CW - 0.6, max(h - 0.5 - h * 0.45, 0.3),
                      f.get("legende", ""), 13, "grey", mini=9, space_after=0)
            y += h + 0.15

    def s_tableau(self, s, sl):
        bottom = self.body_bottom(sl)
        y = BODY_Y
        if sl.get("intro"):
            self.text(s, ML, y, CW, 0.55, sl["intro"], 15, mini=11)
            y += 0.65
        cols = sl.get("colonnes", []) or [""]
        rows = sl.get("lignes", [])[:9]
        com = sl.get("commentaire")
        avail = bottom - y - (0.7 if com else 0)
        cells = [str(c) for r in rows for c in r]
        longest = max([len(c) for c in cells] + [10])
        size = 14 if longest < 40 and len(rows) <= 5 else (12 if len(rows) <= 7 else 11)
        row_h = min(0.6, avail / (len(rows) + 1))
        tbl = s.shapes.add_table(len(rows) + 1, len(cols), Inches(ML), Inches(y), Inches(CW),
                                 Inches(row_h * (len(rows) + 1))).table
        self._fill_table(tbl, cols, rows, size)
        # estimation de la hauteur réelle du tableau
        colw = CW / len(cols)
        real = sum(max(row_h, max(self.need([str(c)], colw, size) for c in (r or [""])) + 0.06) for r in [cols] + rows)
        if real > avail * 1.05:
            self.flag("tableau trop haut pour la slide : réduire le nombre de lignes ou la longueur des cellules")
        if com:
            self.text(s, ML, bottom - 0.65, CW, 0.65, com, 13, "grey", mini=10)

    def s_points(self, s, sl):
        bottom = self.body_bottom(sl)
        enc = sl.get("encadre")
        w = 7.6 if enc else CW
        self.text(s, ML, BODY_Y, w, bottom - BODY_Y, sl.get("points", []), 18, bullets=True, space_after=10, mini=11)
        if enc:
            x = ML + w + 0.35
            ew = W - ML - x
            self.rect(s, x, BODY_Y, ew, bottom - BODY_Y, "red")
            self.text(s, x + 0.3, BODY_Y + 0.3, ew - 0.6, 0.55, enc.get("titre", ""), 16, "white", True, mini=12)
            self.text(s, x + 0.3, BODY_Y + 0.95, ew - 0.6, bottom - BODY_Y - 1.2, enc.get("texte", ""), 15, "white",
                      mini=10)

    def s_chiffres(self, s, sl):
        bottom = self.body_bottom(sl)
        kpis = sl.get("chiffres", [])[:4]
        n = max(len(kpis), 1)
        gap = 0.3
        w = (CW - gap * (n - 1)) / n
        h = 2.2
        for i, k in enumerate(kpis):
            x = ML + i * (w + gap)
            self.rect(s, x, BODY_Y, w, h, "panel")
            self.rect(s, x, BODY_Y + h - 0.07, w, 0.07, "red")
            self.text(s, x + 0.1, BODY_Y + 0.15, w - 0.2, 1.0, k.get("valeur", ""), 40, "red", True,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, mini=20, space_after=0)
            self.text(s, x + 0.2, BODY_Y + 1.2, w - 0.4, 0.85, k.get("libelle", ""), 14, "ink",
                      align=PP_ALIGN.CENTER, mini=10)
        if sl.get("texte"):
            self.text(s, ML, BODY_Y + h + 0.3, CW, bottom - BODY_Y - h - 0.3, sl["texte"], 16, mini=11)

    # -------- code source (fond sombre, chasse fixe, nom du fichier en bandeau)
    def s_code(self, s, sl):
        bottom = self.body_bottom(sl)
        y = BODY_Y
        if sl.get("intro"):
            self.text(s, ML, y, CW, 0.55, sl["intro"], 15, mini=11)
            y += 0.65
        com = sl.get("commentaire")
        cw = CW - 4.0 if com else CW
        lines = (sl.get("code") or "").replace("\t", "    ").rstrip().splitlines() or [""]
        head = 0.36 if sl.get("fichier") else 0
        self.rect(s, ML, y, cw, bottom - y, "dark")
        if head:
            self.rect(s, ML, y, cw, head, "red_dark")
            self.text(s, ML + 0.15, y, cw - 0.3, head, sl["fichier"], 11, "FFFFFF", True, font=self.t["mono"],
                      anchor=MSO_ANCHOR.MIDDLE, auto=False, space_after=0)
        box_h = bottom - y - head - 0.2
        size = 14
        while size > 8:
            cpl = max(10, (cw - 0.4) * 72 / (size * 0.6))
            if sum(max(1, math.ceil(len(l) / cpl)) for l in lines) * size * 1.2 / 72 <= box_h:
                break
            size -= 0.5
        cpl = max(10, (cw - 0.4) * 72 / (size * 0.6))
        if sum(max(1, math.ceil(len(l) / cpl)) for l in lines) * size * 1.2 / 72 > box_h * 1.05:
            self.flag("bloc de code trop long pour la slide : le raccourcir")
        paras = [[(l if l.strip() else " ", {"color": "9FB0B8" if l.lstrip().startswith(("#", "--", "//")) else "EDEAE4"})]
                 for l in lines]
        self.text(s, ML + 0.2, y + head + 0.1, cw - 0.4, box_h, paras, size, font=self.t["mono"], auto=False,
                  space_after=0)
        if com:
            self.side_note(s, ML + cw + 0.3, y, CW - cw - 0.3, bottom - y, "À noter", com)

    # -------- graphiques natifs (éditables dans PowerPoint)
    def s_graphique(self, s, sl):
        g = sl.get("graphique") or {}
        bottom = self.body_bottom(sl)
        y = BODY_Y
        if sl.get("intro"):
            self.text(s, ML, y, CW, 0.55, sl["intro"], 15, mini=11)
            y += 0.65
        com = sl.get("commentaire")
        cw = CW - 4.0 if com else CW
        try:
            self.chart(s, g, ML, y, cw, bottom - y)
        except (ValueError, TypeError, KeyError, IndexError) as e:
            self.flag(f"graphique invalide ({e}) : données à corriger")
            self.text(s, ML, y, cw, 1, "Graphique non généré : données incomplètes.", 14, "grey")
        if com:
            self.side_note(s, ML + cw + 0.3, y, CW - cw - 0.3, bottom - y, "Lecture", com)

    def chart(self, s, g, x, y, w, h):
        kind = g.get("type", "barres")
        cats = [str(c) for c in g.get("categories", [])]
        series = g.get("series") or []
        if not cats or not series:
            raise ValueError("catégories ou séries manquantes")
        unit = g.get("unite") or ""
        cd = CategoryChartData()
        cd.categories = cats

        if kind == "cascade":
            vals = [_num(v) for v in series[0]["valeurs"]][:len(cats)]
            totals = set(g.get("totaux") or [])
            base, up, down, tot, running = [], [], [], [], 0.0
            for i, v in enumerate(vals):
                if i in totals:
                    base.append(0); up.append(0); down.append(0); tot.append(v); running = v
                else:
                    lo = min(running, running + v)
                    base.append(lo); up.append(v if v > 0 else 0); down.append(-v if v < 0 else 0); tot.append(0)
                    running += v
            for name, data in (("base", base), ("Hausse", up), ("Baisse", down), ("Total", tot)):
                cd.add_series(name, data)
            gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
            ch = gf.chart
            colors = [None, self.t["red"], self.t["grey"], self.t["ink"]]
            for ser, col in zip(ch.plots[0].series, colors):
                if col is None:
                    ser.format.fill.background()
                    ser.format.line.fill.background()
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
            ch.plots[0].gap_width = 60
            # étiquettes : la variation (signée) ou le total
            for si, data, sign in ((1, up, ""), (2, down, "−"), (3, tot, "")):
                ser = ch.plots[0].series[si]
                for pi, v in enumerate(data):
                    if v:
                        tf = ser.points[pi].data_label.text_frame
                        tf.text = f"{sign}{v:,.1f}".replace(",", " ").replace(".", ",").replace(",0", "") + (f" {unit}" if unit else "")
                        f = tf.paragraphs[0].runs[0].font
                        f.size, f.bold = Pt(11), True
                        f.color.rgb = RGBColor.from_string("FFFFFF") if v > 0.12 * max(map(abs, vals)) else self.rgb("ink")
                        ser.points[pi].data_label.position = XL_LABEL_POSITION.CENTER
            ch.has_legend = False
        else:
            for ser in series:
                vals = [_num(v) for v in ser["valeurs"]][:len(cats)]
                vals += [None] * (len(cats) - len(vals))
                cd.add_series(str(ser.get("nom", "")), vals)
            ct = {"barres": XL_CHART_TYPE.COLUMN_CLUSTERED, "barres_h": XL_CHART_TYPE.BAR_CLUSTERED,
                  "courbes": XL_CHART_TYPE.LINE_MARKERS, "anneau": XL_CHART_TYPE.DOUGHNUT}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
            gf = s.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
            ch = gf.chart
            plot = ch.plots[0]
            if kind == "anneau":
                for i, pt in enumerate(plot.series[0].points):
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(self.palette[i % len(self.palette)])
                plot.has_data_labels = True
                plot.data_labels.show_percentage = True
                plot.data_labels.show_value = False
                plot.data_labels.font.size = Pt(11)
                ch.has_legend = True
            else:
                for i, ser in enumerate(plot.series):
                    col = RGBColor.from_string(self.palette[i % len(self.palette)])
                    if kind == "courbes":
                        ser.format.line.color.rgb = col
                        ser.format.line.width = Pt(2.25)
                        ser.marker.format.fill.solid()
                        ser.marker.format.fill.fore_color.rgb = col
                        ser.smooth = False
                    else:
                        ser.format.fill.solid()
                        ser.format.fill.fore_color.rgb = col
                if kind != "courbes" and len(cats) * len(series) <= 16:
                    plot.has_data_labels = True
                    plot.data_labels.font.size = Pt(10)
                    plot.data_labels.number_format = "0.0" if any(_num(v) % 1 for ser in series for v in ser["valeurs"]) else "0"
                    plot.data_labels.number_format_is_linked = False
                if kind in ("barres", "barres_h"):
                    plot.gap_width = 70
                ch.has_legend = len(series) > 1
        ch.font.size = Pt(11)
        ch.font.name = self.t["font"]
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if kind != "anneau":
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.rgb("line")
            va.format.line.fill.background()
            if unit:
                va.has_title = True
                va.axis_title.text_frame.text = normalize.txt(unit)
                va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        title = g.get("titre")
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = normalize.txt(title)
            r = ch.chart_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.bold = Pt(13), True
            r.font.color.rgb = self.rgb("ink")
        return gf

    # -------- schémas de flux redessinés
    def s_schema(self, s, sl):
        sc = sl.get("schema") or {}
        nodes = sc.get("noeuds", [])[:7]
        links = sc.get("liens", [])
        bottom = self.body_bottom(sl)
        com = sl.get("commentaire")
        top = BODY_Y
        area_h = bottom - top - (0.75 if com else 0)
        pos = {}
        n = max(len(nodes), 1)
        if sc.get("disposition") == "cycle" and n >= 3:
            bw, bh = min(2.9, CW / 3.4), 1.05
            cx, cy = W / 2, top + area_h / 2
            rx, ry = CW / 2 - bw / 2 - 0.1, area_h / 2 - bh / 2 - 0.05
            for i, nd in enumerate(nodes):
                a = -math.pi / 2 + 2 * math.pi * i / n
                pos[nd["id"]] = (cx + rx * math.cos(a) - bw / 2, cy + ry * math.sin(a) - bh / 2, bw, bh)
        else:
            gap = 0.55
            bw = (CW - gap * (n - 1)) / n
            bh = min(2.0, area_h * 0.6)
            yy = top + (area_h - bh) / 2
            for i, nd in enumerate(nodes):
                pos[nd["id"]] = (ML + i * (bw + gap), yy, bw, bh)
        # flèches d'abord (sous les boîtes)
        for lk in links:
            a, b = pos.get(lk.get("de")), pos.get(lk.get("vers"))
            if not a or not b:
                self.flag(f"schéma : lien vers un nœud inconnu ({lk})")
                continue
            ax, ay = a[0] + a[2] / 2, a[1] + a[3] / 2
            bx, by = b[0] + b[2] / 2, b[1] + b[3] / 2
            dx, dy = bx - ax, by - ay
            # points d'attache sur les bords des boîtes
            def edge(p, ddx, ddy, out=True):
                x0, y0, w0, h0 = p
                cx0, cy0 = x0 + w0 / 2, y0 + h0 / 2
                if abs(ddx) * h0 > abs(ddy) * w0:
                    k = (w0 / 2) / abs(ddx)
                else:
                    k = (h0 / 2) / abs(ddy) if ddy else 0
                sgn = 1 if out else -1
                return cx0 + sgn * ddx * k, cy0 + sgn * ddy * k
            x1, y1 = edge(a, dx, dy, True)
            x2, y2 = edge(b, dx, dy, False)
            self.arrow(s, x1, y1, x2, y2, "grey")
            if lk.get("libelle"):
                mx, my = (x1 + x2) / 2, (y1 + y2) / 2
                lw = min(2.2, max(0.9, len(lk["libelle"]) * 0.09))
                tb = self.text(s, mx - lw / 2, my - 0.17, lw, 0.34, lk["libelle"], 10, "red_dark", True,
                               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, auto=False, space_after=0)
                tb.fill.solid()
                tb.fill.fore_color.rgb = self.rgb("white")
        for nd in nodes:
            x, y, w, h = pos[nd["id"]]
            self.rect(s, x, y, w, h, "panel", MSO_SHAPE.ROUNDED_RECTANGLE, line="red")
            self.text(s, x + 0.1, y + 0.06, w - 0.2, 0.42, nd.get("titre", ""), 14, "red", True,
                      align=PP_ALIGN.CENTER, mini=10, space_after=0)
            self.text(s, x + 0.1, y + 0.45, w - 0.2, h - 0.5, nd.get("texte", ""), 12, "ink",
                      align=PP_ALIGN.CENTER, mini=9, space_after=0)
        if com:
            self.text(s, ML, bottom - 0.7, CW, 0.7, com, 13, "grey", mini=10)

    # ------------------------------------------------------------------ assemblage
    def build(self, path):
        self.cover()
        self.objectives()
        for i, sec in enumerate(self.c["sections"]):
            self.cur = None
            self.divider(i + 1, sec)
            for j, sl in enumerate(sec["slides"]):
                self.cur = (i, j)
                s = self.content_slide(sl)
                self.page_map[self.n] = (i, j)
                getattr(self, "s_" + str(sl.get("type", "points")), self.s_points)(s, sl)
                if sl.get("notes"):
                    s.notes_slide.notes_text_frame.text = normalize.txt(sl["notes"])
        self.cur = None
        self.closing()
        self.glossary()
        self.bibliography()
        self.prs.core_properties.title = self.c["titre"]
        self.prs.core_properties.author = self.t["auteur"] or "IActuaire"
        self.prs.core_properties.last_modified_by = self.t["auteur"] or "IActuaire"
        self.prs.core_properties.comments = self.t["credit"]
        self.prs.save(path)
        return path


def render_with_issues(course: dict, path: str) -> tuple[str, dict, dict]:
    """Rend le deck et renvoie (chemin, problèmes par slide (i, j), numéro de slide -> (i, j))."""
    d = Deck(course)
    d.build(path)
    return path, d.issues, d.page_map


def slide_count(course: dict) -> int:
    """Nombre total de slides que produira le cours (couverture, objectifs, intercalaires, contenu, fin)."""
    import math as _m
    n_content = sum(len(s["slides"]) for s in course["sections"])
    g = _m.ceil(len(course.get("glossaire") or []) / GLOSS_PER_PAGE)
    src = max(1, _m.ceil(len(course.get("bibliographie") or []) / SOURCES_PER_PAGE))
    return 2 + len(course["sections"]) + n_content + 1 + g + src


def fit_to_slides(course: dict, max_slides: int) -> int:
    """Retire des slides (fin de la section la plus longue, puis glossaire) jusqu'à respecter la limite."""
    removed = 0
    while slide_count(course) > max_slides:
        secs = [s for s in course["sections"] if len(s["slides"]) > 1]
        if secs:
            max(secs, key=lambda s: len(s["slides"]))["slides"].pop()
        elif len(course.get("glossaire") or []) > GLOSS_PER_PAGE:
            course["glossaire"] = course["glossaire"][:GLOSS_PER_PAGE]
        elif course.get("glossaire"):
            course["glossaire"] = []
        else:
            break
        removed += 1
    return removed
